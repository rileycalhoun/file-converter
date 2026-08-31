"""Regenerate the repository-owned Office fixtures used for local WASM smoke tests."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


FIXTURES = Path(__file__).resolve().parent

document = Document()
document.add_heading("File Converter", level=1)
document.add_paragraph("Generated fixture for offline DOCX to PDF conversion.")
document.save(FIXTURES / "sample.docx")

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
title = slide.shapes.title
title.text = "File Converter"
textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
textbox.text_frame.text = "Generated fixture for offline PPTX to PDF conversion."
presentation.save(FIXTURES / "sample.pptx")

workbook = Workbook()
sheet = workbook.active
sheet.title = "Offline conversion"
sheet.append(["Format", "Expected output"])
sheet.append(["XLSX", "PDF"])
workbook.save(FIXTURES / "sample.xlsx")

(FIXTURES / "sample.rtf").write_text(
    r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs28 File Converter\par Generated fixture for offline RTF conversion.}",
    encoding="ascii",
)
(FIXTURES / "sample.txt").write_text(
    "File Converter\nGenerated fixture for offline TXT to PDF conversion.\n",
    encoding="utf-8",
)
(FIXTURES / "sample.csv").write_text(
    "format,expected output\nCSV,PDF\n",
    encoding="utf-8",
)

Image.new("RGB", (32, 24), (62, 121, 167)).save(FIXTURES / "sample.jpg", quality=90)
Image.new("RGBA", (32, 24), (125, 86, 160, 255)).save(FIXTURES / "sample.png")
